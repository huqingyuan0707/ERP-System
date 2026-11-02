"""商品导入导出/任务回归；业务用例独立内存 SQLite，不读取或访问业务数据库。"""

import asyncio
from contextlib import asynccontextmanager
from decimal import Decimal
from io import BytesIO

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.v1.endpoints import export as export_api
from app.api.v1.endpoints import task as task_api
from app.core.deps import get_current_tenant, get_current_user, get_db
from app.core.exceptions import AppException, register_exception_handlers
from app.db.base import Base
from app.models.base_data import BdGoods, BdGoodsCategory, BdUnit, BdWarehouse
from app.models.inventory import InvStock
from app.models.sys import SysOperLog, SysTask, SysUser
from app.schemas.goods import GoodsQuery
from app.schemas.task import TaskQuery
from app.services import excel_export
from app.services.auth import AuthService
from app.services.excel_export import ExcelExportService
from app.services.task import TaskService
from app.utils import excel, importer
from app.utils.excel import build_workbook
from app.utils.importer import (
    IMPORT_HEADERS,
    LEGACY_IMPORT_HEADERS,
    TEMPLATE_TIPS,
    build_import_template,
    parse_import_rows,
    to_decimal,
)


@asynccontextmanager
async def memory_db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with async_sessionmaker(engine, expire_on_commit=False)() as db:
            yield db
    finally:
        await engine.dispose()


def content_for(*items):
    return build_workbook(
        IMPORT_HEADERS, ([item.get(header) for header in IMPORT_HEADERS] for item in items)
    )


def test_template_and_exact_legacy_tips():
    content = build_import_template()
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        assert workbook.sheetnames == ["商品导入", "填写说明与示例"]
        assert len(list(workbook.worksheets[0].values)) == 1
    finally:
        workbook.close()
    assert parse_import_rows(content) == []
    legacy = build_workbook(LEGACY_IMPORT_HEADERS, [TEMPLATE_TIPS, ["A", "商品"]])
    assert [row for row, _ in parse_import_rows(legacy)] == [3]
    not_tips = TEMPLATE_TIPS.copy()
    not_tips[0] = "真实编码"
    rows = parse_import_rows(build_workbook(LEGACY_IMPORT_HEADERS, [not_tips]))
    assert rows[0][0] == 2
    assert rows[0][1]["商品编码"] == "真实编码"


@pytest.mark.parametrize("headers", [["商品名称"], ["商品编码", "商品名称", "商品编码"]])
def test_required_and_duplicate_headers(headers):
    with pytest.raises(AppException, match="表头"):
        parse_import_rows(build_workbook(headers, []))


@pytest.mark.parametrize("value", ["no", "NaN", "Infinity", "1,000", True, "=1+1", "--2"])
def test_decimal_rejects_invalid_instead_of_zero(value):
    with pytest.raises(ValueError):
        to_decimal(value)
    assert to_decimal(None) == Decimal(0)
    assert to_decimal("9007199254740991.01") == Decimal("9007199254740991.01")


def test_import_resource_limits_and_invalid_file(monkeypatch):
    with pytest.raises(AppException, match="xlsx"):
        parse_import_rows(b"not an excel file")
    monkeypatch.setattr(importer, "IMPORT_MAX_ROWS", 1)
    with pytest.raises(AppException, match="不能超过"):
        parse_import_rows(
            content_for({"商品编码": "1", "商品名称": "A"}, {"商品编码": "2", "商品名称": "B"})
        )
    monkeypatch.setattr(importer, "IMPORT_MAX_UNCOMPRESSED_BYTES", 10)
    with pytest.raises(AppException, match="资源上限"):
        parse_import_rows(build_import_template())


def test_import_atomic_errors_and_same_goods_rules():
    async def scenario():
        async with memory_db() as db:
            db.add(BdUnit(id=1, tenant_id=1, name="件", precision=0))
            await db.commit()
            result = await ExcelExportService(db, 1).import_goods(
                content_for(
                    {"商品编码": "GOOD", "商品名称": "正常", "成本价": "1.20"},
                    {"商品编码": "INVALID", "商品名称": "错误", "成本价": "oops", "售价": "0.123"},
                    {"商品编码": "UNIT", "商品名称": "精度", "单位名称": "件", "库存上限": "1.5"},
                    {"商品编码": "STATUS", "商品名称": "状态", "状态": "2"},
                    {"商品编码": "X" * 33, "商品名称": "长度"},
                    {"商品编码": "NEG", "商品名称": "负数", "库存下限": "-1"},
                    {"商品编码": "RANGE", "商品名称": "范围", "库存下限": "2", "库存上限": "1"},
                ),
                11,
            )
            assert not result.success and result.success_count == 0
            assert result.fail_count == 6
            assert result.errors[0].row == 3
            assert any("成本价" in error for error in result.errors[0].errors)
            assert any("售价" in error for error in result.errors[0].errors)
            assert (await db.scalar(select(func.count()).select_from(BdGoods))) == 0
            logs = (await db.scalars(select(SysOperLog))).all()
            assert len(logs) == 1 and logs[0].status == 0 and logs[0].tenant_id == 1

    asyncio.run(scenario())


def test_import_tenant_paths_duplicates_and_success_audit():
    async def scenario():
        async with memory_db() as db:
            db.add_all(
                [
                    BdGoodsCategory(id=1, tenant_id=1, name="甲"),
                    BdGoodsCategory(id=2, tenant_id=1, name="乙"),
                    BdGoodsCategory(id=3, tenant_id=1, name="叶", parent_id=1),
                    BdGoodsCategory(id=4, tenant_id=1, name="叶", parent_id=2),
                    BdGoodsCategory(id=5, tenant_id=2, name="外部分类"),
                    BdUnit(id=1, tenant_id=1, name="个", precision=4),
                    BdUnit(id=2, tenant_id=2, name="外部单位", precision=4),
                    BdGoods(tenant_id=2, code="SAME", name="其他租户", barcode="SAME-BC"),
                ]
            )
            await db.commit()
            svc = ExcelExportService(db, 1)
            result = await svc.import_goods(
                content_for(
                    {"商品编码": "A", "商品名称": "A", "分类名称": "叶"},
                    {
                        "商品编码": "B",
                        "商品名称": "B",
                        "分类名称": "外部分类",
                        "单位名称": "外部单位",
                    },
                    {"商品编码": "C", "商品名称": "C", "条码": "DUP"},
                    {"商品编码": "C", "商品名称": "C", "条码": "DUP"},
                ),
                11,
            )
            assert result.fail_count == 3
            assert "歧义" in result.errors[0].errors[0]
            assert any("文件内重复" in error for error in result.errors[-1].errors)
            result = await svc.import_goods(
                content_for(
                    {
                        "商品编码": "SAME",
                        "商品名称": " 本租户 ",
                        "分类名称": "甲/叶",
                        "单位名称": "个",
                        "条码": "SAME-BC",
                        "成本价": "12.34",
                        "库存上限": "1.2345",
                        "状态": "停用",
                    }
                ),
                11,
            )
            assert result.success and result.success_count == 1
            goods = await db.scalar(select(BdGoods).where(BdGoods.tenant_id == 1))
            assert goods.category_id == 3 and goods.unit_id == 1 and goods.status == 0
            assert goods.cost_price == Decimal("12.34") and goods.created_by == 11
            assert goods.name == "本租户"
            log = await db.scalar(select(SysOperLog).where(SysOperLog.status == 1))
            assert log.tenant_id == 1 and log.user_id == 11

    asyncio.run(scenario())


def test_import_audit_failure_rolls_back_all_goods(monkeypatch):
    async def fail_audit(*args, **kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(excel_export, "write_oper_log", fail_audit)

    async def scenario():
        async with memory_db() as db:
            with pytest.raises(RuntimeError, match="audit unavailable"):
                await ExcelExportService(db, 1).import_goods(
                    content_for(
                        {"商品编码": "A", "商品名称": "A"},
                        {"商品编码": "B", "商品名称": "B"},
                    ),
                    11,
                )
            assert await db.scalar(select(func.count()).select_from(BdGoods)) == 0

    asyncio.run(scenario())


def test_export_more_than_10000_and_relations_are_tenant_scoped(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_export, "EXPORT_DIR", tmp_path)

    async def scenario():
        async with memory_db() as db:
            db.add_all(
                [
                    BdGoodsCategory(id=1, tenant_id=1, name="根"),
                    BdGoodsCategory(id=2, tenant_id=1, name="叶", parent_id=1),
                    BdGoodsCategory(id=3, tenant_id=2, name="保密分类"),
                    BdUnit(id=1, tenant_id=2, name="保密单位"),
                ]
            )
            db.add_all(
                [
                    BdGoods(tenant_id=1, code=f"G{i:05d}", name="商品", category_id=2)
                    for i in range(10001)
                ]
            )
            db.add(
                BdGoods(tenant_id=1, code="BROKEN", name="=not-a-formula", category_id=3, unit_id=1)
            )
            db.add(BdGoods(tenant_id=2, code="SECRET", name="保密商品"))
            await db.commit()
            task = await ExcelExportService(db, 1).export_goods(11, GoodsQuery())
            assert task.status == 3 and task.can_download
            workbook = load_workbook(task.file_path, read_only=True)
            try:
                rows = list(workbook.active.values)
                assert len(rows) == 10003
                assert rows[1][3] == "根/叶"
                assert rows[-1][1] == "=not-a-formula"
                assert rows[-1][3] is None and rows[-1][4] is None
                assert "保密" not in str(rows)
            finally:
                workbook.close()
            assert await db.scalar(select(SysTask.tenant_id)) == 1

    asyncio.run(scenario())


def test_export_limit_failure_is_explicit_and_removes_partial_file(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_export, "EXPORT_DIR", tmp_path)
    monkeypatch.setattr(excel, "XLSX_MAX_ROWS", 2)

    async def scenario():
        async with memory_db() as db:
            db.add_all([BdGoods(tenant_id=1, code=str(i), name="商品") for i in range(2)])
            await db.commit()
            task = await ExcelExportService(db, 1).export_goods(11, GoodsQuery())
            assert task.status == 4 and "上限" in task.error_msg
            assert not task.file_path and not task.can_download
            assert not list(tmp_path.glob("*.xlsx"))

    asyncio.run(scenario())


def test_stock_export_filters_stock_goods_warehouse_and_unit(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_export, "EXPORT_DIR", tmp_path)

    async def scenario():
        async with memory_db() as db:
            db.add_all(
                [
                    BdUnit(id=1, tenant_id=2, name="外部单位"),
                    BdGoods(id=1, tenant_id=1, code="A", name="本租户", unit_id=1),
                    BdGoods(id=2, tenant_id=2, code="B", name="外部商品"),
                    BdWarehouse(id=1, tenant_id=1, code="W1", name="本仓库"),
                    BdWarehouse(id=2, tenant_id=2, code="W2", name="外部仓库"),
                    InvStock(tenant_id=1, goods_id=1, warehouse_id=1, qty=2, avg_cost=3),
                    InvStock(tenant_id=1, goods_id=2, warehouse_id=1, qty=2),
                    InvStock(tenant_id=1, goods_id=1, warehouse_id=2, qty=2),
                    InvStock(tenant_id=2, goods_id=1, warehouse_id=1, qty=2),
                ]
            )
            await db.commit()
            task = await ExcelExportService(db, 1).export_stocks(11)
            assert task.status == 3
            workbook = load_workbook(task.file_path, read_only=True)
            try:
                rows = list(workbook.active.values)
                assert len(rows) == 2 and "外部" not in str(rows)
            finally:
                workbook.close()

    asyncio.run(scenario())


def test_task_owner_admin_tenant_and_exact_file_lookup(tmp_path):
    async def scenario():
        async with memory_db() as db:
            path = tmp_path / "a_file.xlsx"
            path.write_bytes(b"fixture")
            db.add_all(
                [
                    SysTask(
                        tenant_id=1,
                        task_id="own",
                        task_type="goods_export",
                        status=3,
                        created_by=11,
                        file_path=str(path),
                    ),
                    SysTask(
                        tenant_id=1,
                        task_id="other",
                        task_type="goods_export",
                        status=3,
                        created_by=12,
                        file_path=str(tmp_path / "other.xlsx"),
                    ),
                    SysTask(
                        tenant_id=2,
                        task_id="foreign",
                        task_type="goods_export",
                        status=3,
                        created_by=11,
                    ),
                    SysTask(
                        tenant_id=None,
                        task_id="legacy",
                        task_type="goods_export",
                        status=3,
                        created_by=11,
                    ),
                ]
            )
            await db.commit()
            svc = TaskService(db, 1)
            own = await svc.list_tasks(TaskQuery(), user_id=11, codes=[])
            assert own.total == 1 and own.items[0].can_download
            all_tasks = await svc.list_tasks(TaskQuery(), user_id=11, codes=["sys:task:view"])
            assert all_tasks.total == 2
            other = await svc.get_task_by_id("other", user_id=11, codes=["sys:task:view"])
            assert not other.can_download and other.file_path is None
            for task_id in ["other", "foreign", "legacy"]:
                with pytest.raises(AppException) as exc:
                    await svc.get_task_by_id(task_id, user_id=11, codes=[])
                assert exc.value.code == 404
            with pytest.raises(AppException) as exc:
                await svc.get_task_by_id("foreign", user_id=11, codes=["*"])
            assert exc.value.code == 404
            for codes in [["sys:task:view"], ["report:export"]]:
                with pytest.raises(AppException) as exc:
                    await svc.download_path(
                        user_id=12, codes=codes, export_dir=tmp_path, task_id="own"
                    )
                assert exc.value.code == 403
            assert (
                await svc.download_path(user_id=11, codes=[], export_dir=tmp_path, task_id="own")
                == path
            )
            assert (
                await svc.download_path(
                    user_id=12, codes=["*"], export_dir=tmp_path, file_name=path.name
                )
                == path
            )
            for name in ["file.xlsx", "%file.xlsx", "a%file.xlsx"]:
                with pytest.raises(AppException) as exc:
                    await svc.download_path(
                        user_id=11, codes=[], export_dir=tmp_path, file_name=name
                    )
                assert exc.value.code == 404
            with pytest.raises(AppException) as exc:
                await svc.download_path(
                    user_id=11, codes=[], export_dir=tmp_path, file_name="../a_file.xlsx"
                )
            assert exc.value.code == 400
            with pytest.raises(ValueError):
                ExcelExportService(db, None)

    asyncio.run(scenario())


def test_http_owner_routes_template_filename_and_xls_rejection(tmp_path, monkeypatch):
    monkeypatch.setattr(task_api, "EXPORT_DIR", tmp_path)

    async def scenario():
        async with memory_db() as db:
            path = tmp_path / "owner.xlsx"
            path.write_bytes(b"fixture")
            db.add(
                SysTask(
                    tenant_id=1,
                    task_id="own",
                    task_type="goods_export",
                    status=3,
                    created_by=11,
                    file_path=str(path),
                )
            )
            await db.commit()
            app = FastAPI()
            app.include_router(export_api.router)
            app.include_router(task_api.router)
            register_exception_handlers(app)
            app.dependency_overrides[get_db] = lambda: db
            app.dependency_overrides[get_current_user] = lambda: SysUser(
                id=11, tenant_id=1, username="owner", password_hash="unused"
            )
            app.dependency_overrides[get_current_tenant] = lambda: 1

            async def permissions(self, user_id):
                return ["base:goods:import", "base:goods:export"]

            monkeypatch.setattr(AuthService, "get_permission_codes", permissions)
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://test"
            ) as client:
                assert (await client.get("/sys/tasks")).status_code == 200
                assert (await client.get("/sys/tasks/own")).json()["data"]["canDownload"]
                assert (await client.get("/sys/tasks/own/download")).content == b"fixture"
                template = await client.get("/base/goods/template")
                assert template.status_code == 200
                assert "filename*=UTF-8''%" in template.headers["content-disposition"]
                assert parse_import_rows(template.content) == []
                response = await client.post(
                    "/base/goods/import", files={"file": ("old.xls", b"old")}
                )
                assert response.status_code == 400 and "另存" in response.json()["message"]

    asyncio.run(scenario())
