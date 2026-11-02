"""商品导入模板与严格的 xlsx 解析；保留真实 Excel 行号。"""

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.core.exceptions import AppException, ErrorCode
from app.utils.excel import ExcelWriter


LEGACY_IMPORT_HEADERS = [
    "商品编码",
    "商品名称",
    "规格",
    "分类名称",
    "单位名称",
    "条码",
    "成本价",
    "售价",
    "库存下限",
    "库存上限",
    "备注",
]
IMPORT_HEADERS = [*LEGACY_IMPORT_HEADERS, "状态"]
REQUIRED_HEADERS = {"商品编码", "商品名称"}
TEMPLATE_TIPS = ["", "", "", "需先在系统中创建", "需先在系统中创建", "", "", "", "", "", ""]
IMPORT_MAX_ROWS = 5000
IMPORT_MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
IMPORT_MAX_COLUMNS = 64
DECIMAL_PATTERN = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z")


def build_import_template() -> bytes:
    """数据页只放表头；说明及示例单独放置，空模板不能导入示例商品。"""
    writer = ExcelWriter(IMPORT_HEADERS, "商品导入")
    try:
        instructions = writer.workbook.create_sheet("填写说明与示例")
        instructions.append(["请在“商品导入”页填写数据；本页不会被导入。"])
        instructions.append(["仅支持 .xlsx；商品编码、商品名称必填，全部通过后整批入库。"])
        instructions.append(["分类/单位须已维护；分类填写完整路径（如 食品/饮料）或无歧义的单名。"])
        instructions.append(["金额最多2位小数，库存上下限最多4位；不可为负数，空白数值默认为0。"])
        instructions.append(
            ["编码与条码请设置为文本以保留前导零；状态为启用/停用或1/0，空白默认启用。"]
        )
        instructions.append(IMPORT_HEADERS)
        instructions.append(
            [
                "GD0001",
                "示例商品",
                "规格A",
                "食品/饮料",
                "个",
                "6901234567890",
                "10.00",
                "20.00",
                "5",
                "100",
                "仅示例，请勿复制为真实数据",
                "启用",
            ]
        )
        buffer = BytesIO()
        writer.save(buffer)
        return buffer.getvalue()
    finally:
        writer.close()


def parse_import_rows(content: bytes) -> list[tuple[int, dict[str, object]]]:
    """验证结构并按数据页解析，限行限列限解压大小，拒绝公式及非法文件。"""
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > IMPORT_MAX_UNCOMPRESSED_BYTES:
                raise ValueError("Excel 解压后超过 64MB 资源上限")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise AppException(
            ErrorCode.VALIDATION_ERROR, f"无法读取 Excel，请使用有效 .xlsx 文件：{exc}"
        ) from exc
    try:
        sheet = (
            workbook["商品导入"] if "商品导入" in workbook.sheetnames else workbook.worksheets[0]
        )
        if sheet.max_column and sheet.max_column > IMPORT_MAX_COLUMNS:
            raise AppException(ErrorCode.VALIDATION_ERROR, "导入列数超过64列，请使用商品导入模板")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        header_index = _map_headers(header)
        result: list[tuple[int, dict[str, object]]] = []
        for excel_row, row in enumerate(rows, start=2):
            # 只兼容旧模板的第2行、原11列表头以及完整说明内容，绝不宽泛跳过错误行。
            if (
                excel_row == 2
                and [to_str(cell) for cell in header] == LEGACY_IMPORT_HEADERS
                and [to_str(cell) for cell in row] == TEMPLATE_TIPS
            ):
                continue
            if all(to_str(value) == "" for value in row):
                continue
            item = {
                name: row[idx] if idx < len(row) else None for name, idx in header_index.items()
            }
            result.append((excel_row, item))
            if len(result) > IMPORT_MAX_ROWS:
                raise AppException(
                    ErrorCode.VALIDATION_ERROR, f"单次导入不能超过 {IMPORT_MAX_ROWS} 行"
                )
        return result
    finally:
        workbook.close()


def _map_headers(header_row: tuple[object, ...]) -> dict[str, int]:
    """必需列不可缺失，任何非空表头都不可重复，以免覆盖列导致静默丢数据。"""
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = to_str(cell)
        if not name:
            continue
        if name in mapping:
            raise AppException(ErrorCode.VALIDATION_ERROR, f"第1行表头重复：{name}")
        mapping[name] = idx
    missing = REQUIRED_HEADERS - mapping.keys()
    if missing:
        raise AppException(
            ErrorCode.VALIDATION_ERROR, f"第1行缺少必需表头：{'、'.join(sorted(missing))}"
        )
    return mapping


def to_str(value: object) -> str:
    """单元格转字符串（None -> ''）。"""
    return "" if value is None else str(value).strip()


def to_decimal(value: object, default: Decimal = Decimal(0)) -> Decimal:
    """仅空值使用默认值；非法、非有限数值、布尔值和公式都报错，不降级为0。"""
    text = to_str(value)
    if not text:
        return default
    if isinstance(value, bool) or len(text) > 128 or not DECIMAL_PATTERN.fullmatch(text):
        raise ValueError("必须填写合法数值（不支持千分位分隔符或公式）")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("必须填写合法数值") from exc
    if not number.is_finite():
        raise ValueError("必须填写有限数值")
    return number
