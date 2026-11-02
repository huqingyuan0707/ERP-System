"""Excel 写入工具：逐行写入、显式格式限制，并将字符串按文本保存。"""

from collections.abc import Iterable
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter


XLSX_MAX_ROWS = 1_048_576
XLSX_MAX_COLUMNS = 16_384
XLSX_MAX_CELL_LENGTH = 32_767
EXPORT_MAX_BYTES = 256 * 1024 * 1024


class ExcelWriter:
    """write-only 工作簿，避免把全部导出行和单元格常驻内存。"""

    def __init__(self, headers: list[str], sheet_title: str = "Sheet1") -> None:
        self.workbook = Workbook(write_only=True)
        self.sheet = self.workbook.create_sheet(title=sheet_title)
        self.row_count = 0
        for index, header in enumerate(headers, start=1):
            self.sheet.column_dimensions[get_column_letter(index)].width = max(12, len(header) * 2)
        self.append(headers)

    def append(self, row: Iterable[object]) -> None:
        if self.row_count >= XLSX_MAX_ROWS:
            raise ValueError(f"导出超过 Excel 单页 {XLSX_MAX_ROWS - 1} 条数据上限，请缩小筛选范围")
        cells = []
        for value in row:
            if len(cells) >= XLSX_MAX_COLUMNS:
                raise ValueError("导出列数超过 Excel 格式上限")
            if isinstance(value, str):
                if len(value) > XLSX_MAX_CELL_LENGTH:
                    raise ValueError("导出单元格文本超过 Excel 格式上限")
                cell = WriteOnlyCell(self.sheet, value=value)
                # 名称/编码以 = 开头也必须是文本，不能变成公式。
                cell.data_type = "s"
                cells.append(cell)
            else:
                cells.append(value)
        self.sheet.append(cells)
        self.row_count += 1

    def save(self, target: Path | BytesIO) -> None:
        self.workbook.save(target)
        size = target.stat().st_size if isinstance(target, Path) else target.tell()
        if size > EXPORT_MAX_BYTES:
            raise ValueError("导出文件超过 256MB 资源上限，请缩小筛选范围")

    def close(self) -> None:
        # write-only 未保存而发生异常时也关闭并清理临时 XML。
        for sheet in self.workbook.worksheets:
            if not sheet.closed:
                sheet.close()
            writer = sheet._writer
            if writer is not None and Path(writer.out).exists():
                writer.cleanup()
        self.workbook.close()


def build_workbook(
    headers: list[str], rows: Iterable[Iterable[object]], sheet_title: str = "Sheet1"
) -> bytes:
    """小型模板/文档转为字节；大数据导出应使用 ExcelWriter 直接写文件。"""
    writer = ExcelWriter(headers, sheet_title)
    try:
        for row in rows:
            writer.append(row)
        buffer = BytesIO()
        writer.save(buffer)
        return buffer.getvalue()
    finally:
        writer.close()
